"""Write analysis results to a new Excel file.

Layout:
- Sheet "数据" — original input columns + appended output columns + per-row error column.
- Sheet "汇总" — whole-table summary (Markdown text).
- Sheet "运行信息" — preset/model metadata.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kdv.llm.schema import FieldSpec


HEADER_FILL = PatternFill("solid", fgColor="5B6CFF")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ERROR_FILL = PatternFill("solid", fgColor="FEE2E2")


def _stringify(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        import json

        return json.dumps(v, ensure_ascii=False)
    return v


def write_results(
    *,
    output_path: str | Path,
    input_columns: list[str],
    rows: list[dict[str, Any]],
    output_fields: list[FieldSpec],
    row_outputs: list[dict[str, Any] | None],
    row_errors: list[str | None],
    summary_markdown: str | None = None,
    meta: dict[str, Any] | None = None,
) -> None:
    """Write a results workbook.

    `rows[i]`, `row_outputs[i]`, `row_errors[i]` must be aligned by index.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "数据"

    output_col_names = [f.name for f in output_fields]
    headers = list(input_columns) + output_col_names + ["错误信息"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, row in enumerate(rows):
        out = row_outputs[i] if i < len(row_outputs) else None
        err = row_errors[i] if i < len(row_errors) else None
        line = [_stringify(row.get(c)) for c in input_columns]
        line += [_stringify((out or {}).get(name)) for name in output_col_names]
        line.append(err or "")
        ws.append(line)
        if err:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col_idx).fill = ERROR_FILL

    _autosize(ws, headers)
    ws.freeze_panes = "A2"

    if summary_markdown is not None:
        s = wb.create_sheet("汇总")
        s.append(["整表汇总分析（Markdown）"])
        s.cell(row=1, column=1).fill = HEADER_FILL
        s.cell(row=1, column=1).font = HEADER_FONT
        for line in summary_markdown.splitlines():
            s.append([line])
        s.column_dimensions["A"].width = 100

    info = wb.create_sheet("运行信息")
    info.append(["键", "值"])
    info.cell(row=1, column=1).fill = HEADER_FILL
    info.cell(row=1, column=1).font = HEADER_FONT
    info.cell(row=1, column=2).fill = HEADER_FILL
    info.cell(row=1, column=2).font = HEADER_FONT
    info.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    info.append(["输入行数", len(rows)])
    success = sum(1 for o, e in zip(row_outputs, row_errors) if o and not e)
    info.append(["成功行数", success])
    info.append(["失败行数", len(rows) - success])
    if meta:
        for k, v in meta.items():
            info.append([str(k), _stringify(v)])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 60

    wb.save(output_path)


def _autosize(ws, headers: Iterable[str]) -> None:
    for i, h in enumerate(headers, start=1):
        col_letter = get_column_letter(i)
        max_len = len(str(h))
        for row in ws.iter_rows(min_col=i, max_col=i, min_row=2, values_only=True):
            v = row[0]
            if v is None:
                continue
            s = str(v)
            if len(s) > max_len:
                max_len = len(s)
            if max_len > 50:
                max_len = 50
                break
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 50)

"""Parse an Excel "output template" file into a list of FieldSpec.

Two supported layouts in one workbook:

1. **Explicit field-definition sheet** (preferred). One of these sheet names:
   `字段定义`, `schema`, `Schema`, `metadata`, `Metadata`, `字段`, `输出字段`.
   Recognized columns (in any order, case-insensitive, both Chinese and English):
     - 字段名 / name              (required)
     - 类型 / type                (string|number|integer|boolean|enum|array|date)
     - 描述 / description / 说明
     - 必填 / required            (是/否, true/false, 1/0)
     - 示例 / example / 样例
     - 枚举值 / enum / enum_values  (separated by comma/分号/、/|)
     - 提示词 / prompt / hint     (extra hint appended to description)

2. **Sample-only sheet**. If no field-definition sheet is found, the first sheet
   is treated as filled example rows; types are inferred from the values.

You can also place both a sample sheet and a definition sheet in the same workbook.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from kdv.excel.inferer import infer_fields_from_sample
from kdv.excel.reader import ExcelTable, read_excel
from kdv.llm.schema import FieldSpec


SCHEMA_SHEET_NAMES = {
    s.lower()
    for s in ("字段定义", "schema", "metadata", "字段", "输出字段", "field_definitions", "field definitions")
}

NAME_KEYS = {"字段名", "字段", "name", "field", "column"}
TYPE_KEYS = {"类型", "type", "field_type"}
DESC_KEYS = {"描述", "说明", "description", "desc"}
REQ_KEYS = {"必填", "required", "is_required"}
EXAMPLE_KEYS = {"示例", "样例", "example", "sample"}
ENUM_KEYS = {"枚举值", "enum", "enum_values", "enums", "values"}
PROMPT_KEYS = {"提示词", "prompt", "hint"}

SUPPORTED_TYPES = {"string", "number", "integer", "boolean", "enum", "array", "date"}
TYPE_ALIASES = {
    "字符串": "string",
    "文本": "string",
    "text": "string",
    "str": "string",
    "数字": "number",
    "数值": "number",
    "float": "number",
    "double": "number",
    "整数": "integer",
    "int": "integer",
    "布尔": "boolean",
    "bool": "boolean",
    "日期": "date",
    "时间": "date",
    "datetime": "date",
    "枚举": "enum",
    "选择": "enum",
    "数组": "array",
    "列表": "array",
    "list": "array",
    "multi": "array",
}


@dataclass
class TemplateLoadResult:
    fields: list[FieldSpec]
    source: str  # "explicit" or "inferred"
    sample_columns: list[str]
    sample_rows: list[dict[str, Any]]
    sheet_used: str


def _normalize_key(s: Any) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", "", str(s)).strip().lower()


def _truthy(v: Any) -> bool:
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v != 0
    s = str(v).strip().lower()
    return s in {"true", "yes", "1", "y", "是", "必填", "✓"}


def _parse_enum(v: Any) -> list[str]:
    if v is None:
        return []
    s = str(v).strip()
    if not s:
        return []
    parts = re.split(r"[,，;；、|/]\s*", s)
    return [p.strip() for p in parts if p.strip()]


def _normalize_type(v: Any) -> str:
    if v is None:
        return "string"
    s = str(v).strip().lower()
    if s in SUPPORTED_TYPES:
        return s
    return TYPE_ALIASES.get(s, "string")


def _find_schema_sheet(wb) -> str | None:
    for name in wb.sheetnames:
        if name.strip().lower() in SCHEMA_SHEET_NAMES:
            return name
    return None


def _read_sheet_as_dicts(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = rows[0]
    keys = [_normalize_key(h) for h in header]
    out: list[dict[str, Any]] = []
    for r in rows[1:]:
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in r):
            continue
        d: dict[str, Any] = {}
        for k, v in zip(keys, r):
            if k:
                d[k] = v
        out.append(d)
    return out


def _pick(d: dict[str, Any], keys: set[str]) -> Any:
    norm = {k.lower(): k for k in d.keys()}
    for k in keys:
        kk = k.lower()
        if kk in norm:
            return d[norm[kk]]
    return None


def _parse_explicit_definition(records: list[dict[str, Any]]) -> list[FieldSpec]:
    fields: list[FieldSpec] = []
    for rec in records:
        name = _pick(rec, NAME_KEYS)
        if not name:
            continue
        name = str(name).strip()
        if not name:
            continue
        ftype = _normalize_type(_pick(rec, TYPE_KEYS))
        desc = _pick(rec, DESC_KEYS)
        prompt = _pick(rec, PROMPT_KEYS)
        full_desc_parts = [str(x).strip() for x in (desc, prompt) if x is not None and str(x).strip()]
        full_desc = "；".join(full_desc_parts)
        example = _pick(rec, EXAMPLE_KEYS)
        enum_vals = _parse_enum(_pick(rec, ENUM_KEYS))
        required = _pick(rec, REQ_KEYS)
        required_bool = _truthy(required) if required is not None else True
        if ftype != "enum" and enum_vals:
            ftype = "enum"
        fields.append(
            FieldSpec(
                name=name,
                type=ftype,  # type: ignore[arg-type]
                description=full_desc,
                required=required_bool,
                enum_values=enum_vals,
                example="" if example is None else str(example)[:80],
            )
        )
    return fields


def load_output_template(path: str | Path) -> TemplateLoadResult:
    """Load a template Excel and return field specs + a sample preview.

    Resolution order:
      1. If the workbook has a schema sheet, use it.
      2. Else infer from the first sheet's sample rows.
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        schema_sheet = _find_schema_sheet(wb)
        sample_columns: list[str] = []
        sample_rows: list[dict[str, Any]] = []

        # Always pull a sample preview from the first non-schema sheet
        sample_sheet_name: str | None = None
        for n in wb.sheetnames:
            if n != schema_sheet:
                sample_sheet_name = n
                break

        if sample_sheet_name:
            sample = read_excel(path, sheet_name=sample_sheet_name)
            sample_columns = sample.columns
            sample_rows = sample.head(5)

        if schema_sheet:
            ws = wb[schema_sheet]
            records = _read_sheet_as_dicts(ws)
            fields = _parse_explicit_definition(records)
            if fields:
                return TemplateLoadResult(
                    fields=fields,
                    source="explicit",
                    sample_columns=sample_columns,
                    sample_rows=sample_rows,
                    sheet_used=schema_sheet,
                )

        # Fall back to inference
        if sample_sheet_name:
            sample = ExcelTable(
                columns=sample_columns,
                rows=sample_rows or [],
                sheet_name=sample_sheet_name,
                source_path=str(path),
            )
            fields = infer_fields_from_sample(sample)
            return TemplateLoadResult(
                fields=fields,
                source="inferred",
                sample_columns=sample_columns,
                sample_rows=sample_rows,
                sheet_used=sample_sheet_name,
            )

        return TemplateLoadResult(
            fields=[], source="inferred", sample_columns=[], sample_rows=[], sheet_used=""
        )
    finally:
        wb.close()


def export_template_skeleton(path: str | Path, fields: list[FieldSpec]) -> None:
    """Write a starter template workbook with both a sample sheet and a schema sheet.

    Useful for "下载示例模板" CTA on the model page.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    sample_ws = wb.active
    sample_ws.title = "示例数据"
    sample_ws.append([f.name for f in fields])
    if fields:
        sample_ws.append([f.example for f in fields])

    schema_ws = wb.create_sheet("字段定义")
    headers = ["字段名", "类型", "描述", "必填", "示例", "枚举值", "提示词"]
    schema_ws.append(headers)
    fill = PatternFill("solid", fgColor="5B6CFF")
    font = Font(bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        cell = schema_ws.cell(row=1, column=col_idx)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for f in fields:
        schema_ws.append(
            [
                f.name,
                f.type,
                f.description,
                "是" if f.required else "否",
                f.example,
                "、".join(f.enum_values),
                "",
            ]
        )
    for col_letter, width in zip("ABCDEFG", (16, 10, 32, 8, 18, 24, 32)):
        schema_ws.column_dimensions[col_letter].width = width

    wb.save(Path(path))

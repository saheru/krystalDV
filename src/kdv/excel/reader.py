"""Read Excel files into normalized in-memory tables.

A "table" is one sheet from one .xlsx file. A "workbook" is the union of all
tables the user wants to analyse together — possibly spanning multiple files
and/or multiple sheets per file. The agent and analysis layers consume this
list directly; the UI lets the user load several files and toggle individual
sheets.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl


# Synthetic column name added by `concat_tables` to mark each row's source.
TABLE_ID_COLUMN = "_table_id"


@dataclass
class ExcelTable:
    columns: list[str]
    rows: list[dict[str, Any]]
    sheet_name: str
    source_path: str
    # Human-readable unique id, e.g. "invoices.xlsx/Sheet1". Used by the
    # agent's per-tool `table` argument to disambiguate when multiple tables
    # are loaded; populated by `read_workbook` / `make_table_id`.
    table_id: str = ""

    def __post_init__(self) -> None:
        if not self.table_id:
            self.table_id = make_table_id(self.source_path, self.sheet_name)

    def __len__(self) -> int:
        return len(self.rows)

    def head(self, n: int = 5) -> list[dict[str, Any]]:
        return self.rows[:n]

    def column_values(self, name: str) -> list[Any]:
        return [r.get(name) for r in self.rows]


def make_table_id(source_path: str | Path, sheet_name: str) -> str:
    """Build a stable, human-readable table id from file + sheet."""
    name = Path(source_path).name if source_path else "(untitled)"
    return f"{name}/{sheet_name}" if sheet_name else name


def _normalize_cell(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return v


def _read_sheet(ws, source_path: str) -> ExcelTable:
    rows_iter = ws.iter_rows(values_only=True)
    header: list[str] = []
    for row in rows_iter:
        if any(c is not None and str(c).strip() != "" for c in row):
            header = [
                str(c).strip() if c is not None and str(c).strip() else f"列{i + 1}"
                for i, c in enumerate(row)
            ]
            break
    if not header:
        return ExcelTable(columns=[], rows=[], sheet_name=ws.title, source_path=source_path)

    # de-duplicate column names
    seen: dict[str, int] = {}
    unique_header: list[str] = []
    for name in header:
        if name in seen:
            seen[name] += 1
            unique_header.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 1
            unique_header.append(name)

    rows: list[dict[str, Any]] = []
    for row in rows_iter:
        values = list(row) + [None] * max(0, len(unique_header) - len(row))
        values = values[: len(unique_header)]
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in values):
            continue
        rows.append({c: _normalize_cell(v) for c, v in zip(unique_header, values)})

    return ExcelTable(
        columns=unique_header,
        rows=rows,
        sheet_name=ws.title,
        source_path=source_path,
    )


def read_excel(path: str | Path, sheet_name: str | None = None) -> ExcelTable:
    """Read one sheet (first by default, or named) into ExcelTable.

    Kept for back-compat with callers that intentionally analyze only one
    sheet (e.g. the schema-template loader). For multi-sheet / multi-file
    use cases call `read_workbook` instead.
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
        return _read_sheet(ws, str(path))
    finally:
        wb.close()


def read_workbook(
    paths: Iterable[str | Path],
    *,
    skip_empty: bool = True,
) -> list[ExcelTable]:
    """Read every sheet from every file into a flat list of `ExcelTable`s.

    Sheet order is preserved, file order matches the iterable order. Empty
    sheets (no header row) are skipped by default — they only confuse the
    agent and waste UI space. `table_id` values are guaranteed unique within
    the returned list.
    """
    tables: list[ExcelTable] = []
    seen_ids: set[str] = set()
    for p in paths:
        path = Path(p)
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                t = _read_sheet(ws, str(path))
                if skip_empty and not t.columns:
                    continue
                # Disambiguate if (very rare) two files have same name + sheet.
                base = t.table_id
                tid = base
                i = 2
                while tid in seen_ids:
                    tid = f"{base}#{i}"
                    i += 1
                t.table_id = tid
                seen_ids.add(tid)
                tables.append(t)
        finally:
            wb.close()
    return tables


def concat_tables(
    tables: list[ExcelTable],
    *,
    table_id_column: str = TABLE_ID_COLUMN,
) -> ExcelTable:
    """Stack rows from every table into one virtual table.

    Used by non-agent modes (逐行 / 整表汇总 / 二者都做) which are inherently
    single-table. Adds a `_table_id` column so the LLM can still see which
    rows came from which sheet/file. Columns are unioned: rows missing a
    column get None for that column.

    For a single-table workbook this is a thin wrapper that just adds the
    table_id column — preserving full back-compat with existing analysis.
    """
    if not tables:
        return ExcelTable(
            columns=[], rows=[], sheet_name="", source_path="",
            table_id="(empty)",
        )
    # Union of columns, in first-seen order. _table_id goes first so it's
    # visible in any preview rendering.
    seen: set[str] = set()
    merged_cols: list[str] = [table_id_column]
    seen.add(table_id_column)
    for t in tables:
        for c in t.columns:
            if c not in seen:
                merged_cols.append(c)
                seen.add(c)

    merged_rows: list[dict[str, Any]] = []
    for t in tables:
        for r in t.rows:
            new_row: dict[str, Any] = {table_id_column: t.table_id}
            for c in merged_cols:
                if c == table_id_column:
                    continue
                new_row[c] = r.get(c)
            merged_rows.append(new_row)

    if len(tables) == 1:
        # Single-table case — preserve the original sheet / source so the
        # result page header still says the right thing.
        return ExcelTable(
            columns=merged_cols,
            rows=merged_rows,
            sheet_name=tables[0].sheet_name,
            source_path=tables[0].source_path,
            table_id=tables[0].table_id,
        )
    src_paths = sorted({Path(t.source_path).name for t in tables if t.source_path})
    return ExcelTable(
        columns=merged_cols,
        rows=merged_rows,
        sheet_name=f"合并 {len(tables)} 张表",
        source_path=" + ".join(src_paths) or "",
        table_id="__concat__",
    )

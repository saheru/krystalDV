from pathlib import Path

import openpyxl

from kdv.excel.writer import write_results
from kdv.llm.schema import FieldSpec


def test_write_results_creates_three_sheets(tmp_path: Path) -> None:
    out = tmp_path / "result.xlsx"
    fields = [FieldSpec(name="标签"), FieldSpec(name="评分", type="integer")]
    rows = [{"客户": "Alice", "反馈": "good"}, {"客户": "Bob", "反馈": "bad"}]
    outputs = [{"标签": "好评", "评分": 5}, None]
    errors = [None, "rate limit"]
    write_results(
        output_path=out,
        input_columns=["客户", "反馈"],
        rows=rows,
        output_fields=fields,
        row_outputs=outputs,
        row_errors=errors,
        summary_markdown="# 报告\n\n关键洞察。",
        meta={"preset": "p", "mode": "both"},
    )
    assert out.exists()
    wb = openpyxl.load_workbook(out, data_only=True)
    assert {"数据", "汇总", "运行信息"} <= set(wb.sheetnames)
    data = wb["数据"]
    headers = [c.value for c in next(data.iter_rows(min_row=1, max_row=1))]
    assert headers[:2] == ["客户", "反馈"]
    assert headers[-1] == "错误信息"
    rows_v = list(data.iter_rows(min_row=2, values_only=True))
    assert rows_v[0][2] == "好评"
    assert rows_v[1][-1] == "rate limit"
